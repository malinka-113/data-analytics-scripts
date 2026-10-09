#!/usr/bin/env python
# coding: utf-8

# In[1]:


import re
import pandas as pd
import docx
import os
from pathlib import Path


# ============================================================
# КОНСТАНТЫ
# ============================================================
ALLOWED_SRC_TABLES = {
    'ODS_SCT_AXDP_CDD_TGT',
    'ODS_SCT_AXDP_CDD_SP_SPER_TGT',
    'ODS_SCT_AXDP_CDD_A_ANY1_TGT',  # ← ДОБАВЛЕНА
    'ODS_SCT_AXDP_CDD_A_ANY1_RB_CDDAARRC_TGT',
}
TBL_OVUL = 'ACC_DEAL_ADDITIONAL'
TBL_OVUL_DESC = 'ДОПОЛНИТЕЛЬНЫЕ АТРИБУТЫ СДЕЛОК'


# ============================================================
# 1. Извлечение SQL из Word
# ============================================================
def extract_sql_from_word(doc_path: str) -> str:
    doc = docx.Document(doc_path)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if ('SELECT' in text.upper() or 'WITH' in text.upper()) and len(text) > 100:
                    return re.sub(r'\s+', ' ', text)
    return None


# ============================================================
# 2. Расширенный парсинг SQL с дифференциацией источников
# ============================================================
def parse_sql_with_source(sql_text: str):
    """
    Извлекает атрибуты с указанием источника: SELECT / JOIN / WHERE
    """
    tables = {}
    attrs = []

    # --- Таблицы из FROM и JOIN ---
    for m in re.finditer(
        r'(?:FROM|JOIN)\s+([A-Za-z0-9_.]+)(?:\s+(?:AS\s+)?([A-Za-z0-9_]+))?',
        sql_text, re.IGNORECASE
    ):
        full_name = m.group(1)
        alias = (m.group(2) or full_name).lower()
        tables[alias] = full_name

    # --- Разделяем SQL на блоки ---
    # Находим SELECT ... FROM
    select_match = re.search(
        r'(?:SELECT\s+(?:UNIQUE\s+|DISTINCT\s+)?)(.*?)\s+FROM\s+',
        sql_text, re.IGNORECASE | re.DOTALL
    )
    select_block = select_match.group(1) if select_match else ''

    # Находим все JOIN ... ON ...
    join_blocks = re.findall(
        r'JOIN\s+[A-Za-z0-9_.]+\s+(?:AS\s+)?[A-Za-z0-9_]+\s+ON\s+(.*?)(?=\s+(?:LEFT|RIGHT|INNER|OUTER|CROSS|FULL|JOIN|WHERE|GROUP|ORDER|HAVING|LIMIT|$))',
        sql_text, re.IGNORECASE | re.DOTALL
    )

    # Находим WHERE блок
    where_match = re.search(
        r'WHERE\s+(.*?)(?=\s+(?:GROUP|ORDER|HAVING|LIMIT|$))',
        sql_text, re.IGNORECASE | re.DOTALL
    )
    where_block = where_match.group(1) if where_match else ''

    # --- Извлекаем алиасы из SELECT ---
    select_aliases = {}
    if select_block:
        exprs = _split_by_comma(select_block)
        for expr in exprs:
            as_match = re.search(r'\bAS\s+([A-Za-z_][A-Za-z0-9_]*)\s*$', expr, re.IGNORECASE)
            if as_match:
                alias = as_match.group(1)
                # Находим первую колонку в выражении
                col_refs = re.findall(
                    r'([A-Za-z_][A-Za-z0-9_]*)\.([A-Za-z_][A-Za-z0-9_]*)',
                    expr, re.IGNORECASE
                )
                if col_refs:
                    select_aliases[alias.upper()] = {
                        'table_alias': col_refs[0][0].lower(),
                        'column': col_refs[0][1]
                    }

    # --- Обработка SELECT ---
    if select_block:
        for alias, info in select_aliases.items():
            src_table = tables.get(info['table_alias'], info['table_alias'])
            attrs.append({
                'src_table': src_table,
                'src_col': info['column'],
                'alias': alias,
                'source': 'SELECT'
            })

    # --- Обработка JOIN ON ---
    for join_block in join_blocks:
        col_refs = re.findall(
            r'([A-Za-z_][A-Za-z0-9_]*)\.([A-Za-z_][A-Za-z0-9_]*)',
            join_block, re.IGNORECASE
        )
        for tbl_alias, col_name in col_refs:
            tbl_alias_lower = tbl_alias.lower()
            if tbl_alias_lower in ('NVL', 'UPPER', 'LOWER', 'COALESCE'):
                continue
            src_table = tables.get(tbl_alias_lower, tbl_alias)
            attrs.append({
                'src_table': src_table,
                'src_col': col_name,
                'alias': tbl_alias,
                'source': 'JOIN'
            })

    # --- Обработка WHERE ---
    if where_block:
        col_refs = re.findall(
            r'([A-Za-z_][A-Za-z0-9_]*)\.([A-Za-z_][A-Za-z0-9_]*)',
            where_block, re.IGNORECASE
        )
        for tbl_alias, col_name in col_refs:
            tbl_alias_lower = tbl_alias.lower()
            if tbl_alias_lower in ('NVL', 'UPPER', 'LOWER', 'COALESCE'):
                continue
            src_table = tables.get(tbl_alias_lower, tbl_alias)
            attrs.append({
                'src_table': src_table,
                'src_col': col_name,
                'alias': tbl_alias,
                'source': 'WHERE'
            })

    # --- Дедупликация ---
    seen = set()
    unique_attrs = []
    for a in attrs:
        key = (a['src_table'], a['src_col'], a['source'])
        if key not in seen:
            seen.add(key)
            unique_attrs.append(a)

    return tables, unique_attrs


def _split_by_comma(block: str) -> list:
    """Разбивает блок по запятым с учётом скобок"""
    exprs, cur, depth = [], [], 0
    for ch in block:
        if ch == '(':
            depth += 1; cur.append(ch)
        elif ch == ')':
            depth -= 1; cur.append(ch)
        elif ch == ',' and depth == 0:
            exprs.append(''.join(cur)); cur = []
        else:
            cur.append(ch)
    if cur:
        exprs.append(''.join(cur))
    return exprs


# ============================================================
# 3. Автоматическое извлечение правил маппинга из Word
# ============================================================
def build_ovul_mapping_from_doc(doc_path: str) -> dict:
    """
    Парсит таблицу атрибутов и строит словарь правил маппинга
    """
    doc = docx.Document(doc_path)
    mapping = {}

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) < 9:
                continue

            full_name = cells[1]   # П31
            description = cells[3] # П33
            field_source = cells[7] # П37

            if not full_name or not field_source:
                continue

            if 'П31' in full_name or 'Наименование атрибута' in full_name:
                continue
            if 'Реализовано в' not in full_name:
                continue

            # Извлекаем имя атрибута ОВЮЛ
            ovul_attr = ''
            attr_match = re.search(
                r'Реализовано\s+в\s+колонке:\s*[A-Za-z0-9_]+\.([A-Za-z0-9_]+)',
                full_name, re.IGNORECASE
            )
            if attr_match:
                ovul_attr = attr_match.group(1)

            # Извлекаем описание
            desc_match = re.match(r'^(.+?)\s+Реализовано', full_name, re.IGNORECASE)
            short_desc = desc_match.group(1).strip() if desc_match else ''
            full_desc = f"{TBL_OVUL_DESC}.{short_desc}" if short_desc else TBL_OVUL_DESC

            # Разбиваем поля из П37
            fields = [f.strip().upper() for f in field_source.split(',') if f.strip()]
            for field in fields:
                if field:
                    mapping[field] = {
                        'ovul_attr': ovul_attr,
                        'description': full_desc,
                    }

    return mapping


# ============================================================
# 4. Применение правил маппинга
# ============================================================
def build_result(attrs, ovul_mapping):
    """
    Фильтрует атрибуты и применяет правила маппинга
    """
    rows = []
    for a in attrs:
        # Фильтр по разрешённым таблицам
        if a['src_table'] not in ALLOWED_SRC_TABLES:
            continue

        # Определяем ключ сопоставления
        match_key = (a['alias'] if a['source'] == 'SELECT' else a['src_col']).upper()

        if match_key in ovul_mapping:
            rule = ovul_mapping[match_key]
            ovul_attr = rule['ovul_attr']
            ovul_desc = rule['description']
            select_ovul = ovul_attr
        else:
            ovul_attr = 'в наборе данных'
            ovul_desc = ''
            select_ovul = 'в наборе данных'

        rows.append({
            'Источник': a['source'],
            'Таблица источника': a['src_table'],
            'Атрибут': a['alias'] if a['source'] == 'SELECT' else a['src_col'],
            'Исходная колонка': a['src_col'],
            'Таблица ОВЮЛ': TBL_OVUL,
            'Описание таблицы': TBL_OVUL_DESC,
            'Атрибут ОВЮЛ': ovul_attr,
            'Описание атрибута': ovul_desc,
            'Выбрать атрибут ОВЮЛ': select_ovul,
        })

    # Убираем дубликаты
    seen = set()
    unique_rows = []
    for r in rows:
        key = (r['Источник'], r['Таблица источника'], r['Атрибут'])
        if key not in seen:
            seen.add(key)
            unique_rows.append(r)

    return unique_rows


# ============================================================
# 5. Главная функция
# ============================================================
def main():
    word_file = "MAP_E020214_AXDP_ДОПОЛНИТЕЛЬНЫЕ АТРИБУТЫ СДЕЛОК.docx"
    output_file = word_file.replace('.docx', '_FULL_RESULT.xlsx')

    if not os.path.exists(word_file):
        print(f"❌ Файл не найден: {word_file}")
        return

    # 1. Извлекаем SQL
    sql_text = extract_sql_from_word(word_file)
    if not sql_text:
        print("❌ SQL не найден")
        return
    print(f"✅ SQL найден ({len(sql_text)} симв.)")

    # 2. Парсим с дифференциацией
    tables, attrs = parse_sql_with_source(sql_text)
    print(f"✅ Найдено таблиц: {len(tables)}, атрибутов: {len(attrs)}")

    # Выводим статистику по источникам
    from_select = sum(1 for a in attrs if a['source'] == 'SELECT')
    from_join = sum(1 for a in attrs if a['source'] == 'JOIN')
    from_where = sum(1 for a in attrs if a['source'] == 'WHERE')
    print(f"   SELECT: {from_select}, JOIN: {from_join}, WHERE: {from_where}")

    # 3. Извлекаем правила маппинга
    ovul_mapping = build_ovul_mapping_from_doc(word_file)
    print(f"✅ Правил маппинга: {len(ovul_mapping)}")

    # 4. Применяем правила
    rows = build_result(attrs, ovul_mapping)
    print(f"✅ После фильтрации: {len(rows)} строк")

    # 5. Сохраняем
    df = pd.DataFrame(rows)
    cols_order = [
        'Источник', 'Таблица источника', 'Атрибут', 'Исходная колонка',
        'Таблица ОВЮЛ', 'Описание таблицы', 'Атрибут ОВЮЛ',
        'Описание атрибута', 'Выбрать атрибут ОВЮЛ'
    ]
    df = df[[c for c in cols_order if c in df.columns]]
    df.to_excel(output_file, index=False)

    print(f"\n✅ Результат сохранён: {output_file}")
    print("\n" + "=" * 120)
    print(df.to_string(index=False))


if __name__ == "__main__":
    main()


# In[ ]:




